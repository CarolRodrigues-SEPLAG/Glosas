from pathlib import Path
import io
from openpyxl import load_workbook

import pandas as pd

from app import (
    normalize_motivo,
    officialize_motivo,
    parse_qrp_bytes_to_records,
    drop_duplicate_glosa_records,
    consolidate_records,
    export_records_excel,
    clean_motivo_text,
)


def parse_unique(path):
    records = parse_qrp_bytes_to_records(path.read_bytes(), path.name)
    df = pd.DataFrame(records)
    return df.drop_duplicates(subset=['Hospital', 'AIH', 'Valor_Glosa'], keep='first')


def test_normalize_reviewed_motivos():
    cases = {
        'PROFISSIONAL AUTÔNOMO NÃO CADASTRADO (DOC )': 'PROFISSIONAL AUTONOMO NAO CADASTRADO',
        'PROFISSIONAL AUTÔNOMO NÃO CADASTRADO NO HOSPITAL': 'PROFISSIONAL AUTONOMO NAO CADASTRADO NO HOSPITAL',
        'PROFISSIONAL AUTÔNOMO NÃO CADASTRADO NO HOSPITAL COM CBO INFORMADO': 'PROFISSIONAL AUTONOMO NAO CADASTRADO NO HOSPITAL',
        'NÚMERO DA AIH FORA DE FAIXA': 'NUMERO DA AIH FORA DE FAIXA',
        'DÍGITO VERIFICADOR AIH ANTERIOR INVÁLIDO': 'DIGITO VERIFICADOR AIH ANTERIOR INVALIDO',
        'AIH BLOQUEADA POR DUPL.REINTERNAÇÃO MESMO CID DIAS': 'AIH BLOQUEADA POR DUPL.REINTERNACAO, MESMO CID< 3 DIAS',
        'AIH BLOQUEADA POR A PEDIDO/ÓBITO/TRANSFERÊNCIA/EVASÃO C/ DIA P/PROCED. C/MP DIAS ATEND': 'AIH BLOQUEADA POR ALTA A PEDIDO/OBITO/TRANSFERENCIA/EVASAO C/ 1 DIA',
        'LANÇAMENTO OBRIGATÓRIO DE OPM. VERIFIQUE COMPATIBILIDADE NO SIGTAP': 'LANCAMENTO OBRIGATORIO DE OPM',
        'AIH REJEITADA NA IMPORTAÇÃO. VERIFIQUE PROTOCOLO': 'AIH REJEITADA NA IMPORTACAO',
        'TOTAL DE DIÁRIAS SUPERIOR AO PERÍODO DE INTERNAÇÃO NA INFORMADA': 'TOTAL DE DIARIAS SUPERIOR AO PERIODO DE INTERNACAO NA COMPETENCIA',
        'AIH BLOQUEADA POR PERMANÊNCIA A MENOR INJUSTIFICADAD': 'AIH BLOQUEADA POR PERMANENCIA A MENOR INJUSTIFICADA',
        'AIH BLOQUEADA POR PERMANÊNCIA A MENOR INJUSTIFICADAI': 'AIH BLOQUEADA POR PERMANENCIA A MENOR INJUSTIFICADA',
    }
    for raw, expected in cases.items():
        motivo, recognized = officialize_motivo(normalize_motivo(raw))
        assert recognized
        assert motivo == expected


def test_aih_inside_motivo_does_not_split_record():
    path = Path('glosas - 2026.03 (JAN)/Arquivos QRP/OSS Nossa Senhora das Gracas.QRP')
    df = parse_unique(path)
    row = df[df['AIH'].eq('2625107095920')]
    assert len(row) == 1
    assert row.iloc[0]['Motivo_Glosa'] == 'NUMERO DA AIH FORA DE FAIXA'
    assert row.iloc[0]['Valor_Glosa'] == 1942.04


def test_eduardo_campos_reviewed_profissional_value_is_separate():
    path = Path('glosas - 2026.04 (FEV)/Arquivos QRP/OSS Eduardo Campos.QRP')
    df = parse_unique(path)
    row = df[df['Valor_Glosa'].eq(41.38)]
    assert len(row) == 1
    assert row.iloc[0]['Motivo_Glosa'] == 'PROFISSIONAL AUTONOMO NAO CADASTRADO NO HOSPITAL'


def test_eduardo_campos_execucao_invalida_e_alta_a_pedido_ficam_separados():
    path = Path('glosas - 2026.04 (FEV)/Arquivos QRP/OSS Eduardo Campos.QRP')
    df = parse_unique(path)
    alta = df[df['Motivo_Glosa'].eq('AIH BLOQUEADA POR ALTA A PEDIDO/OBITO/TRANSFERENCIA/EVASAO C/ 1 DIA')]
    execucao = df[df['Motivo_Glosa'].eq('COMPETENCIA DE EXECUCAO INVALIDA')]
    assert round(alta['Valor_Glosa'].sum(), 2) == 199.33
    assert round(execucao['Valor_Glosa'].sum(), 2) == 1171.50


def test_remove_motivos_distintos_quando_aih_e_valor_sao_iguais():
    df = pd.DataFrame([
        {
            'Arquivo': 'teste.QRP',
            'Hospital': 'HOSPITAL TESTE',
            'AIH': '2626106714836',
            'Motivo_Glosa': 'COMPETENCIA DE EXECUCAO INVALIDA',
            'Valor_Glosa': 1171.50,
            'Motivo_Reconhecido': True,
        },
        {
            'Arquivo': 'teste.QRP',
            'Hospital': 'HOSPITAL TESTE',
            'AIH': '2626106714836',
            'Motivo_Glosa': 'QUANTIDADES DIFERENTES DO PROCEDIMENTO NA EQUIPE CIRURGICA',
            'Valor_Glosa': 1171.50,
            'Motivo_Reconhecido': True,
        },
    ])

    df_unique = drop_duplicate_glosa_records(df)

    assert len(df_unique) == 1
    assert df_unique.iloc[0]['Motivo_Glosa'] == 'COMPETENCIA DE EXECUCAO INVALIDA'
    assert df_unique['Valor_Glosa'].sum() == 1171.50


def test_restauracao_mantem_informacoes_e_periodos_sobrepostos_separados():
    path = Path('glosas - 2026.04 (FEV)/Arquivos QRP/6 Restauracao.QRP')
    df = pd.DataFrame(parse_qrp_bytes_to_records(path.read_bytes(), path.name))
    df_unique = drop_duplicate_glosa_records(df)

    informacoes = df_unique[df_unique['Motivo_Glosa'].eq('AIH BLOQUEADA POR INFORMACOES OU REGISTROS INCOMPATIVEIS')]
    sobrepostos = df_unique[df_unique['Motivo_Glosa'].eq('AIH BLOQUEADA POR PERIODOS DE INTERNACAO SOBREPOSTOS NO MOVIMENTO')]

    assert len(df_unique) == 14
    assert round(df_unique['Valor_Glosa'].sum(), 2) == 43109.07
    assert round(informacoes['Valor_Glosa'].sum(), 2) == 10843.55
    assert round(sobrepostos['Valor_Glosa'].sum(), 2) == 457.98


def test_geral_de_areias_mantem_periodo_sobreposto_validado_separado():
    path = Path('glosas - 2026.04 (FEV)/Arquivos QRP/DGAR Geral de Areias.QRP')
    df = pd.DataFrame(parse_qrp_bytes_to_records(path.read_bytes(), path.name))
    df_unique = drop_duplicate_glosa_records(df)

    dupl = df_unique[df_unique['Motivo_Glosa'].eq('AIH BLOQUEADA POR DUPL.INTERNACAO C/INTERSERCCAO DE PERIODOS')]
    sobrepostos = df_unique[df_unique['Motivo_Glosa'].eq('AIH BLOQUEADA POR PERIODOS DE INTERNACAO SOBREPOSTOS NO MOVIMENTO')]

    assert len(df_unique) == 11
    assert round(df_unique['Valor_Glosa'].sum(), 2) == 5738.16
    assert round(dupl['Valor_Glosa'].sum(), 2) == 1238.83
    assert round(sobrepostos['Valor_Glosa'].sum(), 2) == 439.69


def make_qrp(*segments):
    return b'\x00\x00'.join(s.encode('utf-16le') for s in segments)


def test_metadata_and_equivalent_motivos_in_excel():
    raw = make_qrp(
        'Competência: 02/2026', 'CNES : DEFINITIVO',
        'CNES : 0123456 - HOSPITAL TESTE',
        '2626100000001', '0301060070',
        'AIH BLOQUEADA POR PERMANÊNCIA A MENOR INJUSTIFICADAI', '4.249,78',
        '2626100000002', '0301060070',
        'AIH BLOQUEADA POR PERMANENCIA A MENOR INJUSTIFICADA', '1.961,98',
    )
    df = pd.DataFrame(parse_qrp_bytes_to_records(raw, 'teste.QRP'))
    assert df['CNES'].tolist() == ['0123456', '0123456']
    assert df['Competência'].tolist() == ['02/2026', '02/2026']
    consolidated = consolidate_records(drop_duplicate_glosa_records(df))
    assert len(consolidated) == 1
    assert consolidated.iloc[0]['Status'] == 'Oficial'
    assert round(consolidated.iloc[0]['Valor_Glosa'], 2) == 6211.76
    # Verifica também a aba de revisão e os zeros iniciais no Excel.
    unknown = df.iloc[[0]].copy()
    unknown['Motivo_Glosa'] = 'MOTIVO AINDA DESCONHECIDO'
    unknown['Motivo_Reconhecido'] = False
    details = pd.concat([df, unknown], ignore_index=True)
    wb = load_workbook(io.BytesIO(export_records_excel(consolidate_records(details), details)))
    assert len(wb.sheetnames) == 3
    for sheet in wb:
        rows = list(sheet.values)
        assert rows[1][rows[0].index('CNES')] == '0123456'
        assert rows[1][rows[0].index('Competência')] == '02/2026'


def test_dedup_preserves_different_units_and_months():
    base = {'Hospital': 'HOSPITAL TESTE', 'CNES': '0123456', 'Competência': '01/2026',
            'AIH': '2626100000001', 'Motivo_Glosa': 'TESTE',
            'Valor_Glosa': 100.0, 'Motivo_Reconhecido': False}
    df = pd.DataFrame([base, base, {**base, 'Competência': '02/2026'},
                       {**base, 'CNES': '0123457'}])
    unique = drop_duplicate_glosa_records(df)
    assert len(unique) == 3
    assert len(consolidate_records(unique)) == 3
    assert unique['Valor_Glosa'].sum() == 300


def test_missing_metadata_and_distinct_motivo():
    raw = make_qrp('2626100000001', 'MOTIVO DESCONHECIDO', '10,00')
    row = parse_qrp_bytes_to_records(raw, '02-2026.QRP')[0]
    assert row['CNES'] == ''
    assert row['Competência'] == ''
    different = 'AIH BLOQUEADA POR PERMANENCIA A MAIOR INJUSTIFICADA'
    assert normalize_motivo(different) == different


def test_metadata_from_real_qrp():
    path = Path('Exemplo QRP/Barão de Lucena.QRP')
    records = parse_qrp_bytes_to_records(path.read_bytes(), path.name)
    assert records
    assert {r['CNES'] for r in records} == {'2427427'}
    assert {r['Competência'] for r in records} == {'08/2025'}


def test_numeric_codes_are_removed_without_losing_normative_reference():
    cases = {
        'QUANTIDADE DE OPM SUPERIOR AO PERMITIDO (0404030050/0702050482/6)':
            ('QUANTIDADE DE OPM SUPERIOR AO PERMITIDO', True),
        'QUANTIDADE DE OPM SUPERIOR AO PERMITIDO (0404030050/0702050482/6)S':
            ('QUANTIDADE DE OPM SUPERIOR AO PERMITIDO', True),
        'QUANTIDADE DE OPM SUPERIOR AO PERMITIDO ( / / )S':
            ('QUANTIDADE DE OPM SUPERIOR AO PERMITIDO', True),
        'AIH BLOQUEADA POR DUPLICIDADE DE ACORDO COM PT 10 DE 06/01/14(ORTOPEDIA)':
            ('AIH BLOQUEADA POR DUPLICIDADE DE ACORDO COM PT 10 DE 06/01/14(ORTOPEDIA)', False),
    }
    for raw, expected in cases.items():
        assert officialize_motivo(normalize_motivo(clean_motivo_text(raw))) == expected


def test_opm_normalization_for_any_hospital():
    rows = []
    for cnes, hospital in [('0000477', 'HOSPITAL A'), ('0000426', 'HOSPITAL B')]:
        raw = make_qrp('Competência: 07/2026', f'CNES : {cnes} - {hospital}',
                       '2626101472995', 'QUANTIDADE DE OPM SUPERIOR AO PERMITIDO (0404030050/0702050482/6)S',
                       '4.915,21')
        rows.extend(parse_qrp_bytes_to_records(raw, 'teste.QRP'))
    result = consolidate_records(drop_duplicate_glosa_records(pd.DataFrame(rows)))
    assert len(result) == 2
    assert set(result['Status']) == {'Oficial'}
    assert set(result['Motivo_Glosa']) == {'QUANTIDADE DE OPM SUPERIOR AO PERMITIDO'}
